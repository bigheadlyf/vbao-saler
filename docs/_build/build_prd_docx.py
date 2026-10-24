from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'docs' / '出海企业商机发现Agent项目开发需求文档_v1.0.md'
OUTPUT = ROOT / 'docs' / '出海企业商机发现Agent项目开发需求文档_v1.0.docx'
FONT = 'Microsoft YaHei'
doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(1.9)
sec.bottom_margin = Cm(1.9)
sec.left_margin = Cm(2.1)
sec.right_margin = Cm(2.1)
sec.header_distance = Cm(.8)
sec.footer_distance = Cm(.8)

def fonts(style, size, bold=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        rf.set(qn('w:' + key), FONT)
    for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        rf.attrib.pop(qn('w:' + key), None)

for sty in ['Normal', 'Body Text', 'List Bullet', 'List Number']:
    fonts(doc.styles[sty], 10.5)
    pf = doc.styles[sty].paragraph_format
    pf.line_spacing = 1.18
    pf.space_after = Pt(5)
    pf.widow_control = True
fonts(doc.styles['Title'], 23, True)
doc.styles['Title'].paragraph_format.space_after = Pt(10)
for sty, size, before, after in [('Heading 1',15,15,7),('Heading 2',11.5,9,5),('Heading 3',10.5,8,4)]:
    fonts(doc.styles[sty], size, True)
    pf = doc.styles[sty].paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.keep_with_next = True
    pf.keep_together = True
for style in doc.styles:
    if style.type == 1 and (style.name.startswith('TOC') or style.name.startswith('toc')):
        fonts(style, 10)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.line_spacing = 1.0
for border_element in list(doc.styles.element.xpath('.//w:pBdr')):
    border_element.getparent().remove(border_element)

def inline(par, txt):
    # Keep Word content identical to the Markdown source; URLs become clickable.
    parts = re.split(r'(https?://[^\s]+)', txt)
    for part in parts:
        if part.startswith(('https://','http://')):
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), par.part.relate_to(part, RT.HYPERLINK, is_external=True))
            r = OxmlElement('w:r')
            rp = OxmlElement('w:rPr')
            col = OxmlElement('w:color'); col.set(qn('w:val'), '155487'); rp.append(col)
            sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '19'); rp.append(sz)
            r.append(rp)
            t = OxmlElement('w:t'); t.text = part; r.append(t)
            link.append(r); par._p.append(link)
        else:
            par.add_run(part)

def table_widths(headers):
    first = headers[0]
    if len(headers) == 4:
        return [1.5, 3.0, 5.0, 7.3]
    if first == '编号':
        return [1.6, 3.3, 11.9]
    if first == '用户提供的客户名称':
        return [5.0, 4.2, 7.6]
    if first == '实体':
        return [3.6, 7.1, 6.1]
    if first == '维度':
        return [1.9, 6.8, 8.1]
    if first == '等级':
        return [1.7, 10.2, 4.9]
    if first == '工作表':
        return [3.0, 3.5, 10.3]
    if first == '版本':
        return [1.5, 3.5, 11.8]
    if first == '场景':
        return [3.2, 8.0, 5.6]
    return [3.0, 6.3, 7.5]

def make_table(lines):
    rows = [[c.strip() for c in x.strip().strip('|').split('|')] for x in lines]
    rows = [r for r in rows if not all(re.fullmatch(r':?-+:?', c) for c in r)]
    widths = table_widths(rows[0])
    tab = doc.add_table(rows=0, cols=len(rows[0]))
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False
    for c, width in zip(tab.columns, widths): c.width = Cm(width)
    pr = tab._tbl.tblPr
    border = OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9')
        border.append(e)
    pr.append(border)
    for i, data in enumerate(rows):
        row = tab.add_row()
        trpr = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit'); trpr.append(no_split)
        if i == 0:
            repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
        for j, (cell, txt) in enumerate(zip(row.cells, data)):
            cell.width = Cm(widths[j]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = cell._tc.get_or_add_tcPr()
            marg = OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                e = OxmlElement('w:' + side); e.set(qn('w:w'),'75'); e.set(qn('w:type'),'dxa'); marg.append(e)
            cp.append(marg)
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'DBE6F0' if i == 0 else ('F5F7F9' if i % 2 == 0 else 'FFFFFF')); cp.append(shade)
            p = cell.paragraphs[0]
            pf = p.paragraph_format
            pf.line_spacing = 1.1; pf.space_before = Pt(0); pf.space_after = Pt(0)
            pf.keep_with_next = i == 0; pf.keep_together = True
            if j == 0 and len(txt) <= 10: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(txt); r.font.name = FONT; r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0,0,0); r.bold = i == 0
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 0.3
    p.add_run().font.size = Pt(2)

def add_toc():
    p = doc.add_paragraph('内容导航', style='Heading 2')
    # Outline level 9 keeps the navigation heading itself out of the contents.
    op = OxmlElement('w:outlineLvl'); op.set(qn('w:val'), '9'); p._p.get_or_add_pPr().append(op)
    p = doc.add_paragraph()
    r = p.add_run()
    b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin'); r._r.append(b)
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' TOC \\o "1-1" \\h \\z \\u '; r._r.append(instr)
    sp = OxmlElement('w:fldChar'); sp.set(qn('w:fldCharType'), 'separate'); r._r.append(sp)
    r.add_text('内容导航')
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end'); r._r.append(end)

lines = SOURCE.read_text(encoding='utf-8').splitlines()
i = 0
toc_done = False
while i < len(lines):
    line = lines[i]
    if not line.strip(): i += 1; continue
    if line.startswith('# '):
        doc.add_paragraph(line[2:], style='Title'); i += 1; continue
    if line.startswith('## '):
        first_heading = not toc_done
        if first_heading: add_toc(); toc_done = True
        p = doc.add_paragraph(line[3:], style='Heading 1')
        if first_heading: p.paragraph_format.page_break_before = True
        i += 1; continue
    if line.startswith('### '):
        doc.add_paragraph(line[4:], style='Heading 2'); i += 1; continue
    if line.startswith('|'):
        group = []
        while i < len(lines) and lines[i].startswith('|'):
            group.append(lines[i]); i += 1
        make_table(group); continue
    if line.startswith('- '):
        p = doc.add_paragraph(style='List Bullet'); inline(p,line[2:]); i += 1; continue
    para = []
    while i < len(lines) and lines[i].strip() and not lines[i].startswith(('#','|','- ')):
        para.append(lines[i].rstrip()); i += 1
    p = doc.add_paragraph()
    inline(p, '\n'.join(para))

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run('VBLA PRD 001  |  版本 1.0  |  ')
run.font.size = Pt(8); run.font.name = FONT
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
doc.core_properties.title = '出海企业商机发现Agent项目开发需求文档'
doc.core_properties.subject = '跨境支付聚合产品销售线索发现 一期开发需求'
doc.core_properties.author = ''
doc.core_properties.last_modified_by = ''
doc.core_properties.version = '1.0'
doc.core_properties.comments = ''
settings = doc.settings.element
update = OxmlElement('w:updateFields'); update.set(qn('w:val'),'true'); settings.append(update)
for border_element in list(doc.element.xpath('.//w:pBdr')):
    border_element.getparent().remove(border_element)
doc.save(OUTPUT)
print(OUTPUT)
print(f'Paragraphs: {len(doc.paragraphs)}, tables: {len(doc.tables)}')
