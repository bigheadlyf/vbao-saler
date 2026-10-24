from pathlib import Path
from docx import Document
import re
import json

ROOT = Path(__file__).resolve().parents[2]
md = ROOT / 'docs' / '出海企业商机发现Agent项目开发需求文档_v1.0.md'
wp = md.with_suffix('.docx')
source = md.read_text(encoding='utf-8')
doc = Document(wp)
sections = re.findall(r'^## (\d+) ', source, re.M)
assert sections == [str(n) for n in range(1,15)], sections
features = re.findall(r'^### (FR\d{2}) ', source, re.M)
assert features == [f'FR{n:02}' for n in range(1,21)]
tests = re.findall(r'^\| (TC\d{2}) \|', source, re.M)
assert tests == [f'TC{n:02}' for n in range(1,21)]
refs = set()
for line in source.splitlines():
    if re.match(r'^\| TC\d{2} \|',line):
        refs.update(re.findall(r'FR\d{2}',line.split('|')[2]))
assert set(features) <= refs, set(features)-refs
texts = '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
for needle in ['500','香港帝晶燈飾有限公司','东莞市东城品德电子经营部','纳优服飾有限公司','FR20','TC20','TBD09']:
    assert needle in texts, needle
assert len(doc.tables) == 15
for table in doc.tables:
    assert table.rows[0]._tr.xpath('./w:trPr/w:tblHeader')
assert not doc.styles.element.xpath('.//w:pBdr')
assert not doc.element.xpath('.//w:pBdr')
report = {
    'status':'passed', 'chapters':len(sections), 'functional_requirements':len(features),
    'acceptance_cases':len(tests),'all_features_have_test_references':True,
    'tables':len(doc.tables), 'docx_exists':wp.is_file(),'markdown_exists':md.is_file(),
    'visual_review_required_separately':True
}
(ROOT / 'docs' / '_build' / 'content_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report, ensure_ascii=True))
